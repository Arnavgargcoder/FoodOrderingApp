import os
from openpyxl import load_workbook

from database import get_connection


EXCEL_FILE = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "admin_data.xlsx"
)


def sync_admin_data():

    if not os.path.exists(EXCEL_FILE):
        print("admin_data.xlsx not found.")
        return

    workbook = load_workbook(EXCEL_FILE, data_only=True)
    sheet = workbook.active

    connection = get_connection()
    cursor = connection.cursor()

    for row in sheet.iter_rows(min_row=2, values_only=True):

        admin_id, name, email, phone, password = row

        if not name or not email or not phone or not password:
            continue

        # Check whether admin already exists
        cursor.execute(
            "SELECT id FROM admin WHERE email = %s",
            (email,)
        )

        existing_admin = cursor.fetchone()

        if existing_admin:

            # Update existing admin
            cursor.execute(
                """
                UPDATE admin
                SET name = %s,
                    phone = %s,
                    password = %s
                WHERE email = %s
                """,
                (name, phone, password, email)
            )

            print(f"Admin updated: {email}")

        else:

            # Insert new admin
            cursor.execute(
                """
                INSERT INTO admin
                (name, email, phone, password)
                VALUES (%s, %s, %s, %s)
                """,
                (name, email, phone, password)
            )

            print(f"Admin added: {email}")

    connection.commit()

    cursor.close()
    connection.close()

    workbook.close()

    print("Admin Excel synchronization completed.")


if __name__ == "__main__":
    sync_admin_data()
